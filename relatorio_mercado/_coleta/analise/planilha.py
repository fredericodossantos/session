#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
planilha.py -- gera a planilha .xlsx do estudo (openpyxl), com:
  - dados_brutos: um caso por linha (link clicavel + coluna 'incluido')
  - indicadores_por_area: os mesmos indicadores de indicadores.json,
    calculados POR FORMULA do Excel sobre dados_brutos
  - exclusoes: contagem por area e motivo, tambem por formula
  - custos_referencia: referencias.json
  - fontes: URLs + data_acesso

So LE casos.jsonl / indicadores.json / referencias.json.

Uso:
    python3 planilha.py
    python3 planilha.py --casos casos_teste.jsonl --indicadores indicadores_teste.json --saida planilha_teste.xlsx

Se o LibreOffice (soffice --headless) estiver disponivel, o script recalcula
a planilha gerada e confere se os valores batem com indicadores.json
(tolerancia pequena), imprimindo um relatorio de validacao.
"""

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common as C

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

MOTIVOS_EXCLUSAO = ["equipamento_alto", "obra_grande", "fornecimento_puro", "mencao_passagem"]

# Colunas de dados_brutos, na ordem em que serao escritas.
# (nome_coluna, funcao_extratora(caso) -> valor)
def _vencedores_resumo(c):
    vs = c.get("vencedores") or []
    return "; ".join(f"{v.get('razao_social')}:{v.get('porte')}" for v in vs)


def _grupo_uf(c):
    escopo = c.get("escopo")
    if escopo == "GO":
        return "GO"
    if escopo in C.UFS_VIZINHAS:
        return "vizinhos"
    return escopo


COLUNAS = [
    ("id", lambda c: c.get("id")),
    ("link_pncp", lambda c: c.get("link_pncp")),
    ("data_acesso", lambda c: c.get("data_acesso")),
    ("uf", lambda c: c.get("uf")),
    ("grupo_uf", _grupo_uf),
    ("municipio", lambda c: c.get("municipio")),
    ("orgao", lambda c: c.get("orgao")),
    ("esfera", lambda c: c.get("esfera")),
    ("data_publicacao", lambda c: c.get("data_publicacao")),
    ("modalidade", lambda c: c.get("modalidade")),
    ("objeto", lambda c: c.get("objeto")),
    ("srp", lambda c: bool(c.get("srp")) if c.get("srp") is not None else None),
    ("areas", lambda c: ";".join(c.get("areas") or [])),
    ("area_principal", lambda c: c.get("area_principal")),
    ("valor_estimado_total", lambda c: c.get("valor_estimado_total")),
    ("valor_homologado_total", lambda c: c.get("valor_homologado_total")),
    ("estimado_itens_com_resultado", lambda c: c.get("estimado_dos_itens_com_resultado")),
    ("desconto", lambda c: c.get("desconto")),
    ("desconto_suspeito", lambda c: int(bool(c.get("desconto_suspeito")))),
    ("n_itens", lambda c: c["itens"]["n_itens"]),
    ("n_com_resultado", lambda c: c["itens"]["n_com_resultado"]),
    ("n_homologados", lambda c: c["itens"]["n_homologados"]),
    ("n_desertos", lambda c: c["itens"]["n_desertos"]),
    ("n_fracassados", lambda c: c["itens"]["n_fracassados"]),
    ("n_anulados_cancelados", lambda c: c["itens"]["n_anulados_cancelados"]),
    ("houve_deserta_ou_fracassada", lambda c: int(bool(c["itens"]["houve_deserta_ou_fracassada"]))),
    ("exclusiva_total_meepp", lambda c: int(bool(c["me_epp"]["exclusiva_total"]))),
    ("exclusiva_parcial_meepp", lambda c: int(bool(c["me_epp"]["exclusiva_parcial"]))),
    ("cota_meepp", lambda c: int(bool(c["me_epp"]["cota"]))),
    ("aplicou_beneficio_meepp", lambda c: int(bool(c["me_epp"]["aplicou_beneficio_meepp"]))),
    ("n_fornecedores_vencedores", lambda c: c.get("n_fornecedores_vencedores")),
    ("vencedores_resumo", _vencedores_resumo),
    ("n_participantes", lambda c: c.get("n_participantes")),
    ("tipo_contrato", lambda c: c.get("tipo_contrato")),
    ("vigencia_meses", lambda c: c.get("vigencia_meses")),
    ("cesto_aereo", lambda c: int(bool(c["exigencias"]["cesto_aereo"]["valor"]))),
    ("munck_guindauto", lambda c: int(bool(c["exigencias"]["munck_guindauto"]["valor"]))),
    ("aceita_munck_com_cesto", lambda c: int(bool(c["exigencias"]["aceita_munck_com_cesto"]["valor"]))),
    ("pmoc_responsavel_tecnico", lambda c: int(bool(c["exigencias"]["pmoc_responsavel_tecnico"]["valor"]))),
    ("crea_cft", lambda c: int(bool(c["exigencias"]["crea_cft"]["valor"]))),
    ("frota_minima_trecho", lambda c: c["exigencias"]["frota_minima"]["trecho"]),
    ("atestados_trecho", lambda c: c["exigencias"]["atestados"]["trecho"]),
    ("frac_material", lambda c: c.get("frac_material")),
    ("exclusao_motivo", lambda c: c["exclusao"]["motivo"]),  # None -> celula em branco (para COUNTIFS "<>")
    ("exclusao_evidencia", lambda c: c["exclusao"]["evidencia"]),
    ("revisao_manual", lambda c: int(bool(c["exclusao"]["revisao_manual"]))),
    ("incluido", lambda c: 0 if c["exclusao"]["motivo"] else 1),
]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
LINK_FONT = Font(color="0563C1", underline="single")


def col_letter(nome):
    idx = [n for n, _ in COLUNAS].index(nome) + 1
    return get_column_letter(idx)


def escreve_dados_brutos(wb, casos):
    ws = wb.create_sheet("dados_brutos")
    for i, (nome, _) in enumerate(COLUNAS, start=1):
        cell = ws.cell(row=1, column=i, value=nome)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"

    for r, caso in enumerate(casos, start=2):
        for i, (nome, fn) in enumerate(COLUNAS, start=1):
            try:
                valor = fn(caso)
            except Exception:
                valor = None
            cell = ws.cell(row=r, column=i, value=valor)
            if nome == "link_pncp" and valor:
                cell.hyperlink = valor
                cell.font = LINK_FONT
    ws.column_dimensions[col_letter("objeto")].width = 60
    ws.column_dimensions[col_letter("link_pncp")].width = 45
    ws.column_dimensions[col_letter("exclusao_evidencia")].width = 40
    return ws, len(casos)


def _rng(nome, n):
    letra = col_letter(nome)
    return f"dados_brutos!${letra}$2:${letra}${max(n, 1) + 1}"


def _cond_area_grupo(area, grupo, n):
    """Parte comum ('array de 0/1') usada nas formulas AGGREGATE: caso
    pertence a area (substring em 'areas', separadas por ';') E esta no
    grupo de UF (GO ou vizinhos) E esta incluido."""
    areas_r = _rng("areas", n)
    grupo_r = _rng("grupo_uf", n)
    incl_r = _rng("incluido", n)
    return f'ISNUMBER(SEARCH("{area}",{areas_r}))*({grupo_r}="{grupo}")*({incl_r}=1)'


def _n_amostra_expr(valor_col, area, grupo, n):
    val_r = _rng(valor_col, n)
    cond = _cond_area_grupo(area, grupo, n)
    return f"SUMPRODUCT(({cond})*ISNUMBER({val_r}))"


def _aggregate_median(valor_col, area, grupo, n):
    val_r = _rng(valor_col, n)
    cond = _cond_area_grupo(area, grupo, n)
    n_amostra = _n_amostra_expr(valor_col, area, grupo, n)
    inner = f"_xlfn.AGGREGATE(12,6,{val_r}/(({cond})*ISNUMBER({val_r})))"
    return f'=IF({n_amostra}=0,"",{inner})'


def _aggregate_quartil(valor_col, area, grupo, n, k):
    val_r = _rng(valor_col, n)
    cond = _cond_area_grupo(area, grupo, n)
    n_amostra = _n_amostra_expr(valor_col, area, grupo, n)
    inner = f"_xlfn.AGGREGATE(17,6,{val_r}/(({cond})*ISNUMBER({val_r})),{k})"
    return f'=IF({n_amostra}=0,"",{inner})'


def _countifs_pct(bool_col, area, grupo, n):
    areas_r = _rng("areas", n)
    grupo_r = _rng("grupo_uf", n)
    incl_r = _rng("incluido", n)
    bool_r = _rng(bool_col, n)
    num = f'COUNTIFS({areas_r},"*"&"{area}"&"*",{grupo_r},"{grupo}",{incl_r},1,{bool_r},1)'
    den = f'COUNTIFS({areas_r},"*"&"{area}"&"*",{grupo_r},"{grupo}",{incl_r},1)'
    return f'=IF({den}=0,"",{num}/{den})'


def _sumifs(valor_col, area, grupo, n):
    areas_r = _rng("areas", n)
    grupo_r = _rng("grupo_uf", n)
    incl_r = _rng("incluido", n)
    val_r = _rng(valor_col, n)
    n_amostra = _n_amostra_expr(valor_col, area, grupo, n)
    soma = f'SUMIFS({val_r},{areas_r},"*"&"{area}"&"*",{grupo_r},"{grupo}",{incl_r},1)'
    return f'=IF({n_amostra}=0,"",{soma})'


def _countifs_n(area, grupo, n):
    areas_r = _rng("areas", n)
    grupo_r = _rng("grupo_uf", n)
    incl_r = _rng("incluido", n)
    return f'=COUNTIFS({areas_r},"*"&"{area}"&"*",{grupo_r},"{grupo}",{incl_r},1)'


def escreve_indicadores_por_area(wb, n_casos, indicadores):
    ws = wb.create_sheet("indicadores_por_area")
    headers = [
        "area", "grupo_uf", "volume_ano_indice (copiado de indicadores.json)",
        "n_casos_detalhados (formula)", "valor_total_estimado (formula)",
        "ticket_mediano (formula)", "mediana_participantes (formula)",
        "n_amostra_participantes (formula)",
        "desconto_mediana (formula)", "desconto_q1 (formula)", "desconto_q3 (formula)",
        "n_amostra_desconto (formula)",
        "pct_deserta_ou_fracassada (formula)", "pct_exclusiva_meepp (formula)",
        "pct_exclusiva_parcial_meepp (formula)", "pct_cota_meepp (formula)",
        "pct_continuado (formula)",
    ]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"

    row = 2
    for area in C.AREAS:
        for grupo in ("GO", "vizinhos"):
            ind = indicadores["por_area_grupo"][area][grupo]
            ws.cell(row=row, column=1, value=area)
            ws.cell(row=row, column=2, value=grupo)
            ws.cell(row=row, column=3, value=ind["volume_ano"])
            ws.cell(row=row, column=4, value=_countifs_n(area, grupo, n_casos))
            ws.cell(row=row, column=5, value=_sumifs("valor_estimado_total", area, grupo, n_casos))
            ws.cell(row=row, column=6, value=_aggregate_median("valor_estimado_total", area, grupo, n_casos))
            ws.cell(row=row, column=7, value=_aggregate_median("n_participantes", area, grupo, n_casos))
            n_part_r = _rng("n_participantes", n_casos)
            cond = _cond_area_grupo(area, grupo, n_casos)
            ws.cell(row=row, column=8, value=f"=SUMPRODUCT(({cond})*ISNUMBER({n_part_r}))")
            ws.cell(row=row, column=9, value=_aggregate_median("desconto", area, grupo, n_casos))
            ws.cell(row=row, column=10, value=_aggregate_quartil("desconto", area, grupo, n_casos, 1))
            ws.cell(row=row, column=11, value=_aggregate_quartil("desconto", area, grupo, n_casos, 3))
            desc_r = _rng("desconto", n_casos)
            ws.cell(row=row, column=12, value=f"=SUMPRODUCT(({cond})*ISNUMBER({desc_r}))")
            ws.cell(row=row, column=13, value=_countifs_pct("houve_deserta_ou_fracassada", area, grupo, n_casos))
            ws.cell(row=row, column=14, value=_countifs_pct("exclusiva_total_meepp", area, grupo, n_casos))
            ws.cell(row=row, column=15, value=_countifs_pct("exclusiva_parcial_meepp", area, grupo, n_casos))
            ws.cell(row=row, column=16, value=_countifs_pct("cota_meepp", area, grupo, n_casos))
            # pct_continuado precisa de COUNTIFS por texto (tipo_contrato="continuado")
            areas_r = _rng("areas", n_casos)
            grupo_r = _rng("grupo_uf", n_casos)
            incl_r = _rng("incluido", n_casos)
            tipo_r = _rng("tipo_contrato", n_casos)
            num = f'COUNTIFS({areas_r},"*"&"{area}"&"*",{grupo_r},"{grupo}",{incl_r},1,{tipo_r},"continuado")'
            den = f'COUNTIFS({areas_r},"*"&"{area}"&"*",{grupo_r},"{grupo}",{incl_r},1)'
            ws.cell(row=row, column=17, value=f'=IF({den}=0,"",{num}/{den})')
            row += 1
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    return ws


def escreve_exclusoes(wb, n_casos):
    ws = wb.create_sheet("exclusoes")
    headers = ["area"] + MOTIVOS_EXCLUSAO + ["total_excluidos", "revisao_manual", "n_casos_area"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    areas_r = _rng("areas", n_casos)
    motivo_r = _rng("exclusao_motivo", n_casos)
    revisao_r = _rng("revisao_manual", n_casos)

    row = 2
    for area in C.AREAS:
        ws.cell(row=row, column=1, value=area)
        col = 2
        for motivo in MOTIVOS_EXCLUSAO:
            ws.cell(row=row, column=col,
                     value=f'=COUNTIFS({areas_r},"*"&"{area}"&"*",{motivo_r},"{motivo}")')
            col += 1
        ws.cell(row=row, column=col,
                 value=f'=COUNTIFS({areas_r},"*"&"{area}"&"*",{motivo_r},"<>")')
        col += 1
        ws.cell(row=row, column=col,
                 value=f'=COUNTIFS({areas_r},"*"&"{area}"&"*",{revisao_r},1)')
        col += 1
        ws.cell(row=row, column=col,
                 value=f'=COUNTIFS({areas_r},"*"&"{area}"&"*")')
        row += 1
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    return ws


def escreve_custos_referencia(wb, referencias):
    ws = wb.create_sheet("custos_referencia")
    registros = referencias.get("registros", [])
    campos = ["categoria", "item", "valor", "unidade", "fonte_url", "data_acesso",
              "trecho_ou_pagina", "observacao", "confiabilidade"]
    for i, h in enumerate(campos, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r, reg in enumerate(registros, start=2):
        for i, campo in enumerate(campos, start=1):
            cell = ws.cell(row=r, column=i, value=reg.get(campo))
            if campo == "fonte_url" and reg.get(campo):
                cell.hyperlink = reg.get(campo)
                cell.font = LINK_FONT
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["E"].width = 45
    ws.column_dimensions["G"].width = 40
    ws.column_dimensions["H"].width = 40
    return ws


def escreve_fontes(wb, referencias):
    ws = wb.create_sheet("fontes")
    ws.cell(row=1, column=1, value="fonte_url").font = HEADER_FONT
    ws.cell(row=1, column=2, value="data_acesso").font = HEADER_FONT
    ws.cell(row=1, column=3, value="n_registros").font = HEADER_FONT
    for c in range(1, 4):
        ws.cell(row=1, column=c).fill = HEADER_FILL

    vistos = {}
    ordem = []
    for reg in referencias.get("registros", []):
        url = reg.get("fonte_url")
        if not url:
            continue
        if url not in vistos:
            vistos[url] = {"data_acesso": reg.get("data_acesso"), "n": 0}
            ordem.append(url)
        vistos[url]["n"] += 1

    for r, url in enumerate(sorted(ordem), start=2):
        cell = ws.cell(row=r, column=1, value=url)
        cell.hyperlink = url
        cell.font = LINK_FONT
        ws.cell(row=r, column=2, value=vistos[url]["data_acesso"])
        ws.cell(row=r, column=3, value=vistos[url]["n"])
    ws.column_dimensions["A"].width = 70
    return ws


# --------------------------------------------------------------------------
# validacao via LibreOffice headless
# --------------------------------------------------------------------------

def recalcula_com_soffice(xlsx_path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        cmd = [soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, xlsx_path]
        try:
            subprocess.run(cmd, check=True, timeout=180, capture_output=True)
        except Exception as e:
            print(f"[planilha.py] AVISO: falha ao rodar LibreOffice: {e}")
            return None
        saida = os.path.join(tmp, os.path.basename(xlsx_path))
        if not os.path.exists(saida):
            return None
        wb = load_workbook(saida, data_only=True)
        # extrai so a aba de indicadores para comparar (le antes do tempdir sumir)
        dados = {}
        ws = wb["indicadores_por_area"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            area, grupo = row[0], row[1]
            dados[(area, grupo)] = row
        return dados


def valida(dados_recalculados, indicadores):
    if dados_recalculados is None:
        print("[planilha.py] LibreOffice indisponivel -- validacao das formulas pulada "
              "(confira manualmente abrindo o .xlsx).")
        return

    campos_ordem = [
        ("n_casos_detalhados", 3), ("valor_total_estimado", 4), ("ticket_mediano", 5),
        ("mediana_participantes", 6), ("n_amostra_participantes", 7),
        ("desconto_mediana", 8), ("desconto_q1", 9), ("desconto_q3", 10),
        ("n_amostra_desconto", 11), ("pct_deserta_ou_fracassada", 12),
        ("pct_exclusiva_meepp", 13), ("pct_exclusiva_parcial_meepp", 14),
        ("pct_cota_meepp", 15), ("pct_continuado", 16),
    ]
    n_ok = 0
    n_falhou = 0
    for area in C.AREAS:
        for grupo in ("GO", "vizinhos"):
            esperado = indicadores["por_area_grupo"][area][grupo]
            linha = dados_recalculados.get((area, grupo))
            if linha is None:
                print(f"[planilha.py] VALIDACAO: linha nao encontrada para {area}/{grupo}")
                n_falhou += 1
                continue
            for campo, idx in campos_ordem:
                exp = esperado.get(campo)
                got = linha[idx] if idx < len(linha) else None
                if got == "":
                    got = None
                if exp is None and got is None:
                    n_ok += 1
                    continue
                if exp is None or got is None:
                    print(f"[planilha.py] VALIDACAO DIVERGENTE {area}/{grupo}/{campo}: "
                          f"indicadores.json={exp} planilha={got}")
                    n_falhou += 1
                    continue
                try:
                    diff = abs(float(exp) - float(got))
                    tol = max(1e-6, abs(float(exp)) * 1e-6)
                    if diff <= tol:
                        n_ok += 1
                    else:
                        print(f"[planilha.py] VALIDACAO DIVERGENTE {area}/{grupo}/{campo}: "
                              f"indicadores.json={exp} planilha={got} (diff={diff})")
                        n_falhou += 1
                except (TypeError, ValueError):
                    if str(exp) == str(got):
                        n_ok += 1
                    else:
                        print(f"[planilha.py] VALIDACAO DIVERGENTE {area}/{grupo}/{campo}: "
                              f"indicadores.json={exp!r} planilha={got!r}")
                        n_falhou += 1
    print(f"[planilha.py] validacao das formulas: {n_ok} OK, {n_falhou} divergentes")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--casos", default=os.path.join(C.ANALISE_DIR, "casos.jsonl"))
    ap.add_argument("--indicadores", default=os.path.join(C.ANALISE_DIR, "indicadores.json"))
    ap.add_argument("--referencias", default=C.CUSTOS_REFERENCIAS_PATH)
    ap.add_argument("--saida", default=os.path.join(C.ANALISE_DIR, "planilha.xlsx"))
    ap.add_argument("--sem-validacao", action="store_true")
    args = ap.parse_args()

    casos = C.carregar_jsonl(args.casos)
    with open(args.indicadores, "r", encoding="utf-8") as f:
        indicadores = json.load(f)
    referencias = {"registros": []}
    if os.path.exists(args.referencias):
        with open(args.referencias, "r", encoding="utf-8") as f:
            referencias = json.load(f)

    print(f"[planilha.py] casos: {len(casos)} (de {args.casos})")

    wb = Workbook()
    wb.remove(wb.active)

    ws_dados, n_casos = escreve_dados_brutos(wb, casos)
    escreve_indicadores_por_area(wb, n_casos, indicadores)
    escreve_exclusoes(wb, n_casos)
    escreve_custos_referencia(wb, referencias)
    escreve_fontes(wb, referencias)

    wb.save(args.saida)
    print(f"[planilha.py] planilha gravada em {args.saida}")

    if not args.sem_validacao:
        recalculados = recalcula_com_soffice(args.saida)
        valida(recalculados, indicadores)


if __name__ == "__main__":
    main()
