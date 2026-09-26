from __future__ import annotations

import tempfile
import sqlite3
from types import SimpleNamespace
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from monitor_ac.cli import _args
from monitor_ac.config import carregar
from monitor_ac.persistencia import Historico
from monitor_ac.relatorio import gerar_csv, gerar_html, gerar_xlsx, publicar


class TestConfiguracaoECLI(unittest.TestCase):
    def test_configuracao_malformada_amigavel(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = Path(td) / "ruim.yaml"
            cfg.write_text("filtro: [", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Não foi possível ler"):
                carregar(cfg)

    def test_configuracao_tipo_invalido(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = Path(td) / "ruim.json"
            cfg.write_text('{"filtro":{"termos_inclusao":["split"]},"api":{"modalidades":[]}}', encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "modalidades"):
                carregar(cfg)

    def test_esfera_malformada_rejeitada(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = Path(td) / "ruim.json"
            cfg.write_text('{"filtro":{"termos_inclusao":["split"]},"esferas":[{}]}', encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "esferas"):
                carregar(cfg)

    def test_modalidades_invalidas_rejeitadas(self):
        for value in ("a,2", "1,1", "0"):
            with self.subTest(value=value), self.assertRaises(SystemExit) as exc:
                _args(["--modalidades", value])
            self.assertEqual(exc.exception.code, 2)

    def test_caminhos_relativos_partem_da_pasta_da_config(self):
        with tempfile.TemporaryDirectory() as td:
            pasta = Path(td) / "config-dir"
            pasta.mkdir()
            cfg = pasta / "config.yaml"
            cfg.write_text("filtro:\n  termos_inclusao: [split]\nsaida:\n  pasta: resultados\n  banco: h.sqlite\n  logs: logs\n", encoding="utf-8")
            carregada = carregar(cfg)
            self.assertEqual(Path(carregada["saida"]["pasta"]), pasta / "resultados")
            self.assertEqual(Path(carregada["saida"]["banco"]), pasta / "resultados" / "h.sqlite")
            self.assertEqual(Path(carregada["saida"]["logs"]), pasta / "logs")


def registro(**overrides):
    row = {"estado": "nova", "situacao_me_epp": "Subcontratação ME/EPP",
           "data_encerramento": None, "data_abertura": None, "orgao": "Órgão",
           "unidade": "Unidade", "municipio": "Goiânia", "codigo_ibge": "5208707",
           "esfera": "Municipal", "modalidade": "Pregão", "numero_compra": "1",
           "objeto": "=1+1", "valor_estimado": 2, "situacao_compra": "Aberta",
           "contagem_itens": {"itens": 2, "exclusivos": 0, "cota": 0, "subcontratacao": 1},
           "numero_controle": "PNCP/1", "link_pncp": "javascript:alert(1)",
           "sistema_origem": "Origem", "link_origem": "https://example.com/a?x=1&y=2",
           "termos": ["@SUM(A1:A2)"]}
    row.update(overrides)
    return row


class TestRelatorios(unittest.TestCase):
    def test_csv_formula_e_subcontratacao(self):
        with tempfile.TemporaryDirectory() as td:
            path = gerar_csv([registro()], Path(td) / "x.csv")
            text = path.read_text(encoding="utf-8-sig")
            self.assertIn("'=1+1", text)
            self.assertIn("'@SUM(A1:A2)", text)
            self.assertIn("0/0/1/2", text)

    def test_html_restringe_esquema_de_links(self):
        with tempfile.TemporaryDirectory() as td:
            path = gerar_html([registro()], Path(td) / "x.html", {})
            text = path.read_text(encoding="utf-8")
            self.assertNotIn("javascript:", text)
            self.assertIn('href="https://example.com/a?x=1&amp;y=2"', text)

    def test_html_rejeita_urls_com_credenciais_ou_barras_invertidas(self):
        with tempfile.TemporaryDirectory() as td:
            for url in ("https://user@example.com/path", "https://trusted.example\\@evil.test/"):
                path = gerar_html([registro(link_pncp=url, link_origem="")], Path(td) / "x.html", {})
                text = path.read_text(encoding="utf-8")
                self.assertNotIn("href=", text)

    def test_xlsx_tem_cabecalho_filtros_e_valores(self):
        with tempfile.TemporaryDirectory() as td:
            path = gerar_xlsx([registro()], Path(td) / "x.xlsx")
            workbook = load_workbook(path, read_only=False)
            planilha = workbook["Licitações"]
            self.assertEqual(planilha.freeze_panes, "A2")
            self.assertEqual(planilha.auto_filter.ref, planilha.dimensions)
            self.assertEqual(planilha.cell(1, 1).value, "Novo")
            self.assertEqual(planilha.cell(2, 20).value, "@SUM(A1:A2)")
            self.assertEqual(planilha.cell(2, 20).data_type, "s")
            workbook.close()

    def test_publicacao_unica_e_preserva_ultimos_se_falhar_geracao(self):
        with tempfile.TemporaryDirectory() as td:
            pasta = Path(td)
            (pasta / "ultimo.csv").write_text("anterior-csv", encoding="utf-8")
            (pasta / "ultimo.html").write_text("anterior-html", encoding="utf-8")
            primeiro = publicar([registro()], pasta, {})
            segundo = publicar([registro()], pasta, {})
            self.assertNotEqual(primeiro["csv"], segundo["csv"])
            antes_csv = (pasta / "ultimo.csv").read_bytes()
            antes_html = (pasta / "ultimo.html").read_bytes()
            antes_xlsx = (pasta / "ultimo.xlsx").read_bytes() if (pasta / "ultimo.xlsx").exists() else None
            with patch("monitor_ac.relatorio.gerar_html", side_effect=OSError("disco cheio")):
                with self.assertRaises(OSError):
                    publicar([registro()], pasta, {})
            self.assertEqual((pasta / "ultimo.csv").read_bytes(), antes_csv)
            self.assertEqual((pasta / "ultimo.html").read_bytes(), antes_html)
            if antes_xlsx is not None:
                self.assertEqual((pasta / "ultimo.xlsx").read_bytes(), antes_xlsx)

    def test_publicacao_reverte_o_trio_se_falhar_troca_de_ultimo(self):
        with tempfile.TemporaryDirectory() as td:
            pasta = Path(td)
            antigos = {tipo: f"versao anterior {tipo}".encode() for tipo in ("csv", "html", "xlsx")}
            for tipo, conteudo in antigos.items():
                (pasta / f"ultimo.{tipo}").write_bytes(conteudo)
            substituir_original = Path.replace
            falhou = False

            def falhar_na_troca_html(caminho, destino):
                nonlocal falhou
                if Path(destino).name == "ultimo.html" and not falhou:
                    falhou = True
                    raise OSError("falha simulada na substituição")
                return substituir_original(caminho, destino)

            with patch.object(Path, "replace", falhar_na_troca_html):
                with self.assertRaisesRegex(OSError, "falha simulada"):
                    publicar([registro()], pasta, {})
            for tipo, conteudo in antigos.items():
                self.assertEqual((pasta / f"ultimo.{tipo}").read_bytes(), conteudo)

    def test_migracao_falha_sem_deixar_schema_parcial_ou_db_aberto(self):
        with tempfile.TemporaryDirectory() as td:
            caminho = Path(td) / "legado.sqlite"
            with sqlite3.connect(caminho) as con:
                con.execute("CREATE VIEW licitacoes AS SELECT 1 AS numero_controle")
            con.close()
            with self.assertRaises(Exception):
                Historico(caminho)
            # O constructor encerra a conexão e desfaz tabelas criadas antes da falha.
            caminho.rename(Path(td) / "renomeado.sqlite")
            with sqlite3.connect(Path(td) / "renomeado.sqlite") as con:
                tabelas = {linha[0] for linha in con.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'")}
            con.close()
            self.assertNotIn("execucoes", tabelas)

    def test_cli_converte_falha_de_publicacao_e_fecha_recursos(self):
        from monitor_ac import cli

        with tempfile.TemporaryDirectory() as td:
            pasta = Path(td)
            config = carregar(Path(__file__).resolve().parent.parent / "config.yaml")
            config["saida"] = {"pasta": str(pasta / "saida"), "banco": str(pasta / "hist.db"),
                               "logs": str(pasta / "logs")}
            fake_resultado = SimpleNamespace(
                modalidades_com_falha=[], por_modalidade={}, encontradas=0, filtradas=0,
                me_epp=0, exclusivas=0, novas=0, falhas=0, registros=[], no_escopo=0)
            cliente = SimpleNamespace(falhas=[], requisicoes=0, fechado=False)
            cliente.close = lambda: setattr(cliente, "fechado", True)
            with patch.object(cli, "carregar", return_value=config), \
                    patch.object(cli, "ClientePNCP", return_value=cliente), \
                    patch.object(cli, "executar", return_value=fake_resultado), \
                    patch.object(cli, "publicar", side_effect=OSError("disco cheio")):
                self.assertEqual(cli.main(["--config", str(pasta / "config.yaml"), "--modalidades", "6"]), 2)
            self.assertTrue(cliente.fechado)
            # Os handlers e a conexão SQLite são fechados ao sair, inclusive na falha.
            (pasta / "logs" / "monitor.log").rename(pasta / "log-movido.log")
            (pasta / "hist.db").rename(pasta / "hist-movido.db")


if __name__ == "__main__":
    unittest.main()
