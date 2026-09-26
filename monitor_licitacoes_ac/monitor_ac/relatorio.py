"""Geração de relatórios CSV, HTML e XLSX."""
from __future__ import annotations

import csv
import html
import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .filtros import COTA, EXCLUSIVA, PARCIAL

COLUNAS = [
    ("novo", "Novo"),
    ("situacao_me_epp", "Situação ME/EPP"),
    ("data_encerramento", "Data-limite propostas"),
    ("data_abertura", "Início recebimento propostas"),
    ("orgao", "Órgão"),
    ("unidade", "Unidade"),
    ("municipio", "Município"),
    ("codigo_ibge", "Código IBGE"),
    ("esfera", "Esfera"),
    ("modalidade", "Modalidade"),
    ("numero_compra", "Nº compra"),
    ("objeto", "Objeto"),
    ("valor_estimado", "Valor total estimado (R$)"),
    ("situacao_compra", "Situação da compra"),
    ("itens", "Itens (exclusivos/cota/subcontratação/total)"),
    ("numero_controle", "Nº controle PNCP"),
    ("link_pncp", "Link PNCP"),
    ("sistema_origem", "Sistema de origem"),
    ("link_origem", "Link sistema de origem"),
    ("termos", "Termos encontrados"),
]


def ordenar(registros: list[dict]) -> list[dict]:
    return sorted(registros, key=lambda r: (r["data_encerramento"] or datetime.max, r["orgao"]))


def fmt_data(d: datetime | None) -> str:
    return d.strftime("%d/%m/%Y %H:%M") if d else ""


def fmt_moeda(v: Any) -> str:
    if v in (None, ""):
        return ""
    try:
        s = f"{float(v):,.2f}"
    except (TypeError, ValueError):
        return str(v)
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _itens_txt(r: dict) -> str:
    c = r.get("contagem_itens") or {}
    if not c.get("itens"):
        return ""
    return (f"{c.get('exclusivos', 0)}/{c.get('cota', 0)}/"
            f"{c.get('subcontratacao', 0)}/{c.get('itens', 0)}")


def _neutralizar_csv(valor: str) -> str:
    # Excel trata fórmulas mesmo em campos entre aspas. Espaços de controle ou
    # tabs antes do marcador também são ignorados por alguns leitores.
    pos = 0
    while pos < len(valor) and (valor[pos].isspace() or ord(valor[pos]) < 32):
        pos += 1
    inicio = valor[pos:]
    if inicio.startswith(("=", "+", "-", "@")):
        return "'" + valor
    return valor


def _valor_csv(r: dict, chave: str) -> str:
    if chave == "novo":
        return {"nova": "NOVA", "atualizada": "ATUALIZADA"}.get(r.get("estado", ""), "")
    if chave in ("data_encerramento", "data_abertura"):
        return fmt_data(r.get(chave))
    if chave == "valor_estimado":
        return fmt_moeda(r.get(chave))
    if chave == "itens":
        return _itens_txt(r)
    if chave == "termos":
        return ", ".join(r.get("termos") or [])
    return str(r.get(chave) if r.get(chave) is not None else "")


def gerar_csv(registros: list[dict], caminho: Path) -> Path:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig + ';' abre direto no Excel em português.
    with caminho.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([titulo for _, titulo in COLUNAS])
        for r in ordenar(registros):
            w.writerow([_neutralizar_csv(_valor_csv(r, chave)) for chave, _ in COLUNAS])
    return caminho


def _valor_xlsx(r: dict, chave: str) -> Any:
    """Converte um registro para um valor seguro e legível no Excel."""
    if chave == "novo":
        return {"nova": "NOVA", "atualizada": "ATUALIZADA"}.get(r.get("estado", ""), "")
    if chave in ("data_encerramento", "data_abertura"):
        return fmt_data(r.get(chave))
    if chave == "valor_estimado":
        valor = r.get(chave)
        if valor in (None, ""):
            return ""
        try:
            return float(valor)
        except (TypeError, ValueError):
            return str(valor)
    if chave == "itens":
        return _itens_txt(r)
    if chave == "termos":
        return ", ".join(r.get("termos") or [])
    valor = r.get(chave)
    if valor is None:
        return ""
    texto = str(valor)
    pos = 0
    while pos < len(texto) and (texto[pos].isspace() or ord(texto[pos]) < 32):
        pos += 1
    if texto[pos:].startswith(("=", "+", "-", "@")):
        return "'" + texto
    return texto


def gerar_xlsx(registros: list[dict], caminho: Path) -> Path:
    """Gera uma planilha XLSX com filtros de tabela e cabeçalho congelado."""
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Licitações"
    cabecalhos = [titulo for _, titulo in COLUNAS]
    ws.append(cabecalhos)
    for registro in ordenar(registros):
        ws.append([_valor_xlsx(registro, chave) for chave, _ in COLUNAS])

    cabecalho = PatternFill("solid", fgColor="123B5D")
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = cabecalho
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    for coluna in ws.columns:
        letra = coluna[0].column_letter
        maior = max(len(str(celula.value or "")) for celula in coluna)
        ws.column_dimensions[letra].width = min(max(maior + 2, 12), 48)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho


CSS = """
:root{--bg:#f4f6f8;--card:#fff;--tx:#1d2733;--mut:#5b6875;--bd:#dde3e9;--ex:#0f7b3f;--exbg:#e6f6ec;
--cota:#8a5a00;--cotabg:#fff4dc;--par:#1f5fa8;--parbg:#e7f0fb;--novo:#c62828;--link:#0b57d0}
@media (prefers-color-scheme:dark){:root{--bg:#12171d;--card:#1b222b;--tx:#e6edf3;--mut:#9aa7b4;--bd:#2c3642;
--ex:#5fd08f;--exbg:#10301f;--cota:#f0c05a;--cotabg:#33280f;--par:#7fb2f0;--parbg:#132740;--novo:#ff7373;--link:#8ab4f8}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--tx);
font:15px/1.45 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
main{max-width:860px;margin:0 auto;padding:16px}
h1{font-size:1.25rem;margin:0 0 4px}.sub{color:var(--mut);font-size:.85rem;margin:0 0 12px}
.resumo{display:grid;grid-template-columns:repeat(auto-fit,minmax(120px,1fr));gap:8px;margin:0 0 16px}
.resumo div{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:8px 10px}
.resumo b{display:block;font-size:1.3rem}.resumo span{color:var(--mut);font-size:.8rem}
.card{background:var(--card);border:1px solid var(--bd);border-left:6px solid var(--bd);border-radius:10px;
padding:12px 14px;margin:0 0 12px;overflow-wrap:anywhere}
.card.ex{border-left-color:var(--ex);background:linear-gradient(90deg,var(--exbg),var(--card) 40%)}
.card.cota{border-left-color:var(--cota)}.card.par{border-left-color:var(--par)}
.topo{display:flex;flex-wrap:wrap;gap:6px;align-items:center;margin-bottom:6px}
.tag{font-size:.75rem;font-weight:600;padding:2px 8px;border-radius:999px;border:1px solid var(--bd);color:var(--mut)}
.tag.ex{background:var(--ex);color:var(--card);border-color:var(--ex)}
.tag.cota{background:var(--cotabg);color:var(--cota);border-color:var(--cota)}
.tag.par{background:var(--parbg);color:var(--par);border-color:var(--par)}
.tag.novo{background:var(--novo);color:#fff;border-color:var(--novo)}
.prazo{font-weight:700;margin-left:auto}.prazo.urg{color:var(--novo)}
.obj{margin:4px 0 8px}.org{font-weight:600}.mut{color:var(--mut);font-size:.85rem}
dl{display:grid;grid-template-columns:max-content 1fr;gap:2px 10px;margin:6px 0;font-size:.88rem}
dt{color:var(--mut)}dd{margin:0}
.links a{display:inline-block;margin:6px 10px 0 0;color:var(--link);font-weight:600;text-decoration:none}
.vazio{padding:24px;text-align:center;color:var(--mut)}
"""


def _classe(situacao: str) -> str:
    return {EXCLUSIVA: "ex", COTA: "cota", PARCIAL: "par"}.get(situacao, "")


def _url_web(valor: Any) -> str | None:
    if not isinstance(valor, str) or not valor.strip():
        return None
    valor = valor.strip()
    # Navegadores reinterpretam barras invertidas e caracteres de controle
    # durante a navegação, enquanto urlsplit pode tratá-los de outra forma.
    if "\\" in valor or any(ord(char) < 32 or ord(char) == 127 for char in valor):
        return None
    try:
        parsed = urlsplit(valor)
        if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname:
            return None
        if parsed.username is not None or parsed.password is not None:
            return None
        # Acessar .port valida portas malformadas.
        _ = parsed.port
    except ValueError:
        return None
    return valor


def gerar_html(registros: list[dict], caminho: Path, resumo: dict[str, Any]) -> Path:
    e = html.escape
    agora = datetime.now()
    cards = []
    for r in ordenar(registros):
        cls = _classe(r["situacao_me_epp"])
        enc = r.get("data_encerramento")
        urgente = enc is not None and (enc - agora).total_seconds() < 3 * 86400
        rotulo = ("★ " if cls == "ex" else "") + (r["situacao_me_epp"] or "—")
        tags = [f'<span class="tag {cls}">{e(rotulo)}</span>']
        if r.get("estado") == "nova":
            tags.insert(0, '<span class="tag novo">NOVA</span>')
        elif r.get("estado") == "atualizada":
            tags.insert(0, '<span class="tag">atualizada</span>')
        tags.append(f'<span class="tag">{e(r["esfera"])}</span>')
        links = []
        link_pncp = _url_web(r.get("link_pncp"))
        link_origem = _url_web(r.get("link_origem"))
        if link_pncp:
            links.append(f'<a href="{e(link_pncp)}" target="_blank" rel="noopener">Edital no PNCP</a>')
        if link_origem:
            nome = r["sistema_origem"] or "Sistema de origem"
            links.append(f'<a href="{e(link_origem)}" target="_blank" rel="noopener">{e(nome)}</a>')
        detalhes = [
            ("Município", r["municipio"]),
            ("Modalidade", r["modalidade"] + (f" nº {r['numero_compra']}" if r["numero_compra"] else "")),
            ("Valor estimado", ("R$ " + fmt_moeda(r["valor_estimado"])) if fmt_moeda(r["valor_estimado"]) else "não informado"),
            ("Propostas", f"{fmt_data(r['data_abertura'])} até {fmt_data(enc)}"),
            ("Itens", _itens_txt(r) and f"{_itens_txt(r)} (exclusivos/cota/subcontratação/total)"),
            ("Situação", r["situacao_compra"]),
            ("Nº PNCP", r["numero_controle"]),
        ]
        dl = "".join(f"<dt>{e(k)}</dt><dd>{e(str(v))}</dd>" for k, v in detalhes if v)
        cards.append(f"""<article class="card {cls}">
<div class="topo">{''.join(tags)}<span class="prazo{' urg' if urgente else ''}">⏰ {e(fmt_data(enc) or 's/ data')}</span></div>
<div class="org">{e(r['orgao'])}</div>
<div class="mut">{e(r['unidade'])}</div>
<p class="obj">{e(r['objeto'])}</p>
<dl>{dl}</dl>
<div class="links">{''.join(links)}</div>
</article>""")
    blocos_resumo = "".join(
        f"<div><b>{e(str(v))}</b><span>{e(k)}</span></div>" for k, v in [
            ("encontradas (GO)", resumo.get("encontradas", 0)),
            ("no filtro", resumo.get("filtradas", 0)),
            ("com benefício ME/EPP", resumo.get("me_epp", 0)),
            ("exclusivas ME/EPP", resumo.get("exclusivas", 0)),
            ("novas", resumo.get("novas", 0)),
        ])
    aviso = ""
    if resumo.get("falhas"):
        aviso = (f'<p class="mut">⚠ {resumo["falhas"]} requisição(ões) falharam; '
                 "veja o log. O resultado pode estar incompleto.</p>")
    corpo = "".join(cards) or '<p class="vazio">Nenhuma licitação encontrada com os filtros atuais.</p>'
    documento = f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Licitações ar-condicionado GO</title><style>{CSS}</style></head>
<body><main>
<h1>Licitações de ar-condicionado — Goiás</h1>
<p class="sub">Gerado em {e(fmt_data(agora))} · fonte: PNCP · {e(resumo.get('parametros', ''))}</p>
<section class="resumo">{blocos_resumo}</section>{aviso}
{corpo}
</main></body></html>"""
    caminho.parent.mkdir(parents=True, exist_ok=True)
    caminho.write_text(documento, encoding="utf-8")
    return caminho


def publicar(registros: list[dict], pasta: Path, resumo: dict[str, Any]) -> dict[str, Path]:
    """Prepara o trio de relatórios e atualiza os arquivos atuais com rollback."""
    pasta.mkdir(parents=True, exist_ok=True)
    agora = datetime.now()
    base = agora.strftime("%Y%m%d_%H%M%S_%f")
    carimbo = base
    n = 1
    while any((pasta / f"licitacoes_ac_go_{carimbo}.{ext}").exists() for ext in ("csv", "html", "xlsx")):
        carimbo = f"{base}_{n}"
        n += 1
    destinos = {"csv": pasta / f"licitacoes_ac_go_{carimbo}.csv",
                "html": pasta / f"licitacoes_ac_go_{carimbo}.html",
                "xlsx": pasta / f"licitacoes_ac_go_{carimbo}.xlsx"}
    ordem = ("csv", "html", "xlsx")
    temporarios: dict[str, Path] = {}
    ultimos_temporarios: dict[str, Path] = {}
    try:
        for tipo, gerador in (("csv", lambda p: gerar_csv(registros, p)),
                              ("html", lambda p: gerar_html(registros, p, resumo)),
                              ("xlsx", lambda p: gerar_xlsx(registros, p))):
            fd, nome = tempfile.mkstemp(prefix=f".{tipo}_", suffix=".tmp", dir=pasta)
            os.close(fd)
            temp = Path(nome)
            temporarios[tipo] = temp
            gerador(temp)
        # Todos os formatos estão prontos antes de publicar qualquer nome.
        for tipo in ordem:
            temporarios[tipo].replace(destinos[tipo])
        for tipo in ("csv", "html", "xlsx"):
            temporarios[tipo] = destinos[tipo]

        caminhos_ultimos = {tipo: pasta / f"ultimo.{tipo}" for tipo in ordem}
        anteriores = {tipo: caminho.read_bytes() if caminho.is_file() else None
                      for tipo, caminho in caminhos_ultimos.items()}
        # Prepara os três novos ponteiros antes de trocar qualquer arquivo visível.
        for tipo in ordem:
            temp_fd, temp_nome = tempfile.mkstemp(prefix=f".ultimo_{tipo}_", suffix=".tmp", dir=pasta)
            os.close(temp_fd)
            temp_ultimo = Path(temp_nome)
            ultimos_temporarios[tipo] = temp_ultimo
            temp_ultimo.write_bytes(destinos[tipo].read_bytes())
        atualizados: list[str] = []
        try:
            for tipo in ordem:
                ultimos_temporarios[tipo].replace(caminhos_ultimos[tipo])
                atualizados.append(tipo)
        except OSError as erro_publicacao:
            # A troca de vários arquivos não é atômica no Windows. Restaura o
            # estado anterior se uma substituição individual falhar.
            for tipo in reversed(atualizados):
                caminho = caminhos_ultimos[tipo]
                conteudo = anteriores[tipo]
                if conteudo is None:
                    caminho.unlink(missing_ok=True)
                    continue
                fd, nome = tempfile.mkstemp(prefix=f".rollback_{tipo}_", suffix=".tmp", dir=pasta)
                os.close(fd)
                rollback = Path(nome)
                try:
                    rollback.write_bytes(conteudo)
                    rollback.replace(caminho)
                finally:
                    rollback.unlink(missing_ok=True)
            raise erro_publicacao
        return destinos
    finally:
        for temp in (*temporarios.values(), *ultimos_temporarios.values()):
            if temp not in destinos.values():
                temp.unlink(missing_ok=True)
